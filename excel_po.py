from fastapi import FastAPI
from pydantic import BaseModel
from typing import List
from openpyxl import Workbook, load_workbook
import os
import time

app = FastAPI()

FILE_NAME = "po_data.xlsx"

# ----------------------------
# Models
# ----------------------------
class Item(BaseModel):
    item_name: str
    quantity: int
    unit_price: float

class POData(BaseModel):
    customer_name: str
    items: List[Item]


def safe_save(wb, filename, retries=5):
    for i in range(retries):
        try:
            wb.save(filename)
            return True
        except PermissionError:
            print(f"⚠️ File locked, retrying... ({i+1})")
            time.sleep(1)
    raise Exception("❌ File is locked. Close Excel and try again.")

# ----------------------------
# API Endpoint
# ----------------------------
@app.post("/add-po")
def add_po(data: POData):
    try:
        # ✅ Create file if not exists
        if not os.path.exists(FILE_NAME):
            wb = Workbook()
            ws = wb.active
            ws.title = "PO_Data"
            ws.append(["Customer Name", "Item Name", "Quantity", "Unit Price"])
            safe_save(wb, FILE_NAME)

        # ✅ Load workbook
        wb = load_workbook(FILE_NAME)
        ws = wb.active

        # ✅ Add rows
        for item in data.items:
            ws.append([
                data.customer_name,
                item.item_name,
                item.quantity,
                item.unit_price
            ])

        # ✅ Save safely
        safe_save(wb, FILE_NAME)

        return {
            "status": "success",
            "rows_added": len(data.items)
        }

    except Exception as e:
        return {
            "status": "error",
            "message": str(e)
        }

# ----------------------------
# Health Check (optional)
# ----------------------------
@app.get("/")
def home():
    return {"message": "API is running 😭"}